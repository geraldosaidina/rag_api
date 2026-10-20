"""Statistical summaries for the existing Phase 1 and Phase 2 benchmark artefacts.

Does not regenerate answers, retrieval rankings, or LLM judgements.
Missing scores stay missing. They are never replaced by zero.
"""

from __future__ import annotations

import csv
import hashlib
import json
import math
import unicodedata
from pathlib import Path
from typing import Any, Iterable, Sequence

import numpy as np
from scipy.stats import wilcoxon

from evaluation.metrics.retrieval import score_ranking

EXPECTED_IDS = [f"Q{i:03d}" for i in range(1, 29)]
UNSUPPORTED_IDS = ("Q023", "Q024", "Q025", "Q026")
CONFIGS = ("A1", "A2", "A3", "A4", "A5")
COMPARISONS = (("A1", "A2"), ("A1", "A3"), ("A2", "A3"), ("A3", "A4"), ("A3", "A5"))
IR_METRICS = ("precision_at_5", "recall_at_5", "recall_at_10", "mrr")
PAIRED_METRICS = ("mrr", "recall_at_5", "recall_at_10")
HUMAN_DIMENSIONS = (
    "human_faithfulness",
    "human_answer_relevance",
    "human_citation_semantic_support",
    "human_unsupported_handling",
    "human_archive_correctness",
)
AGREEMENT_DIMENSIONS = (
    ("human_faithfulness", "faithfulness"),
    ("human_answer_relevance", "answer_relevance"),
    ("human_citation_semantic_support", "citation_semantic_support"),
)
MASTER_SEED = 20261020
N_BOOT = 10_000
VALIDATED_HUMAN_REVIEW = (
    "evaluation/human_review/revisao_manual_pfc_validada.xlsx"
)
CONFIRMED_ARCHIVE_SCORES = {
    "Q018": 2,
    "Q023": 2,
    "Q024": 2,
    "Q025": 2,
    "Q026": 2,
}

COLUMN_ALIASES = {
    "question_id": ("question_id", "id", "consulta", "questao"),
    "human_faithfulness": ("human_faithfulness", "faithfulness", "fidelidade", "fidelidade_factual"),
    "human_answer_relevance": (
        "human_answer_relevance",
        "answer_relevance",
        "relevancia",
        "relevancia_da_resposta",
    ),
    "human_citation_semantic_support": (
        "human_citation_semantic_support",
        "citation_semantic_support",
        "suporte_semantico",
        "suporte_semantico_citacoes",
    ),
    "human_unsupported_handling": (
        "human_unsupported_handling",
        "unsupported_handling",
        "tratamento_nao_suportadas",
    ),
    "human_archive_correctness": (
        "human_archive_correctness",
        "archive_correctness",
        "correcao_arquivo",
        "correcao_face_ao_arquivo",
        "correcao_factual_arquivo",
    ),
}


class AnalysisError(Exception):
    """Raised when source artefacts do not satisfy the analysis contract."""


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    digest.update(path.read_bytes())
    return digest.hexdigest()


def load_jsonl(path: Path) -> list[dict[str, Any]]:
    rows = [
        json.loads(line)
        for line in path.read_text(encoding="utf-8").splitlines()
        if line.strip()
    ]
    if not rows:
        raise AnalysisError(f"Empty JSONL: {path}")
    return rows


def child_seed(key: str, master: int = MASTER_SEED) -> int:
    raw = hashlib.sha256(f"{master}:{key}".encode("utf-8")).digest()
    return int.from_bytes(raw[:8], "little") % (2**32)


def _finite(values: Iterable[float | None]) -> list[float]:
    numeric: list[float] = []
    for value in values:
        if value is None:
            continue
        number = float(value)
        if math.isnan(number):
            continue
        numeric.append(number)
    return numeric


def describe(values: Iterable[float | None]) -> dict[str, Any]:
    """Sample descriptive statistics. Standard deviation uses ddof=1."""
    numeric = _finite(values)
    n = len(numeric)
    if n == 0:
        return {
            "n": 0,
            "mean": None,
            "std": None,
            "median": None,
            "min": None,
            "max": None,
        }
    array = np.asarray(numeric, dtype=float)
    return {
        "n": n,
        "mean": float(array.mean()),
        "std": float(array.std(ddof=1)) if n >= 2 else None,
        "median": float(np.median(array)),
        "min": float(array.min()),
        "max": float(array.max()),
    }


def latency_describe(values: Iterable[float | None]) -> dict[str, Any]:
    stats = describe(values)
    numeric = _finite(values)
    if not numeric:
        stats.update({"p90": None, "p95": None, "range": None})
        return stats
    array = np.asarray(numeric, dtype=float)
    stats["p90"] = float(np.quantile(array, 0.90, method="linear"))
    stats["p95"] = float(np.quantile(array, 0.95, method="linear"))
    stats["range"] = float(array.max() - array.min())
    return stats


def bootstrap_mean_ci(
    values: Sequence[float],
    *,
    seed_key: str,
    n_boot: int = N_BOOT,
) -> dict[str, Any]:
    """Percentile 95% CI of the mean. Resamples queries with replacement."""
    array = np.asarray(list(values), dtype=float)
    n = int(array.size)
    if n == 0:
        return {"n": 0, "ci95_low": None, "ci95_high": None, "n_boot": n_boot}
    rng = np.random.default_rng(child_seed(seed_key))
    draws = rng.integers(0, n, size=(n_boot, n))
    means = array[draws].mean(axis=1)
    low, high = np.quantile(means, [0.025, 0.975], method="linear")
    return {
        "n": n,
        "ci95_low": float(low),
        "ci95_high": float(high),
        "n_boot": n_boot,
        "seed_key": seed_key,
        "seed": child_seed(seed_key),
    }


def holm_adjust(named_pvalues: Sequence[tuple[str, float]]) -> dict[str, float]:
    """Holm step-down adjusted p-values, capped at 1, forced monotone."""
    m = len(named_pvalues)
    if m == 0:
        return {}
    ordered = sorted(named_pvalues, key=lambda item: (item[1], item[0]))
    adjusted: dict[str, float] = {}
    running = 0.0
    for index, (name, pvalue) in enumerate(ordered):
        raw = min(1.0, (m - index) * float(pvalue))
        running = max(running, raw)
        adjusted[name] = running
    return adjusted


def wilcoxon_paired(diffs: Sequence[float]) -> dict[str, Any]:
    """Two-sided Wilcoxon signed-rank test. Zero differences are ties and are dropped."""
    array = np.asarray(list(diffs), dtype=float)
    n = int(array.size)
    n_ties = int(np.sum(array == 0))
    nonzero = array[array != 0]
    result: dict[str, Any] = {
        "applicable": False,
        "n": n,
        "n_ties": n_ties,
        "n_nonzero": int(nonzero.size),
        "statistic": None,
        "pvalue": None,
        "method": None,
        "zero_method": "wilcox",
        "alternative": "two-sided",
        "reason": None,
    }
    if nonzero.size == 0:
        result["reason"] = "all_differences_are_zero"
        return result
    if nonzero.size < 1:
        result["reason"] = "no_nonzero_differences"
        return result
    try:
        tested = wilcoxon(
            nonzero,
            zero_method="wilcox",
            alternative="two-sided",
            method="auto",
        )
    except ValueError as exc:
        result["reason"] = f"wilcoxon_not_applicable: {exc}"
        return result
    result["applicable"] = True
    result["statistic"] = float(tested.statistic)
    result["pvalue"] = float(tested.pvalue)
    result["method"] = "scipy.stats.wilcoxon method=auto on non-zero differences"
    return result


def quadratic_weighted_kappa(
    human: Sequence[int],
    auto: Sequence[int],
    *,
    labels: Sequence[int] = (0, 1, 2),
) -> float | None:
    """Cohen's kappa with quadratic weights. None when undefined."""
    if len(human) != len(auto) or not human:
        return None
    k = len(labels)
    index = {label: pos for pos, label in enumerate(labels)}
    matrix = np.zeros((k, k), dtype=float)
    for left, right in zip(human, auto, strict=True):
        if left not in index or right not in index:
            return None
        matrix[index[left], index[right]] += 1.0
    total = float(matrix.sum())
    if total == 0:
        return None
    weights = np.zeros((k, k), dtype=float)
    denom = float((k - 1) ** 2)
    for i in range(k):
        for j in range(k):
            weights[i, j] = ((i - j) ** 2) / denom
    observed = float((weights * matrix).sum() / total)
    row = matrix.sum(axis=1)
    col = matrix.sum(axis=0)
    expected_matrix = np.outer(row, col) / total
    expected = float((weights * expected_matrix).sum() / total)
    if expected == 0:
        return None
    return 1.0 - (observed / expected)


def confusion_counts(
    human: Sequence[int],
    auto: Sequence[int],
    *,
    labels: Sequence[int] = (0, 1, 2),
) -> dict[str, int]:
    counts = {f"h{h}_a{a}": 0 for h in labels for a in labels}
    for left, right in zip(human, auto, strict=True):
        counts[f"h{left}_a{right}"] += 1
    return counts


def _norm_header(value: Any) -> str:
    text = unicodedata.normalize("NFKD", str(value or ""))
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    text = text.strip().lower().replace(" ", "_").replace("-", "_")
    return text


def parse_score(value: Any) -> int | None:
    """Accept 0, 1, 2. Blank and NA stay missing. Other values raise."""
    if value is None:
        return None
    if isinstance(value, float) and math.isnan(value):
        return None
    if isinstance(value, str):
        text = value.strip().lower()
        if text in {"", "na", "n/a", "n.a.", "none", "null", "-", "—"}:
            return None
        value = text
    try:
        number = float(value)
    except (TypeError, ValueError) as exc:
        raise AnalysisError(f"Unrecognised human score: {value!r}") from exc
    if number != int(number) or int(number) not in (0, 1, 2):
        raise AnalysisError(f"Human score outside {{0,1,2}}: {value!r}")
    return int(number)


def load_human_review(path: Path) -> tuple[dict[str, dict[str, Any]], dict[str, Any]]:
    from openpyxl import load_workbook

    if not path.is_file():
        raise AnalysisError(f"Human review workbook not found: {path}")
    workbook = load_workbook(path, data_only=True, read_only=True)
    chosen = None
    headers: list[str] = []
    mapping: dict[str, int] = {}
    for sheet in workbook.worksheets:
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            continue
        headers = [_norm_header(cell) for cell in rows[0]]
        mapping = {}
        for canonical, aliases in COLUMN_ALIASES.items():
            for alias in aliases:
                if alias in headers:
                    mapping[canonical] = headers.index(alias)
                    break
        if "question_id" in mapping and len(mapping) >= 2:
            chosen = (sheet.title, rows)
            break
    if chosen is None:
        raise AnalysisError(
            "No worksheet with a question id column and human score columns. "
            f"Sheets={workbook.sheetnames}"
        )
    sheet_name, rows = chosen
    missing = [name for name in HUMAN_DIMENSIONS if name not in mapping]
    if missing:
        raise AnalysisError(
            f"Sheet '{sheet_name}' is missing columns {missing}. Headers={headers}"
        )
    parsed: dict[str, dict[str, Any]] = {}
    for raw in rows[1:]:
        if raw is None or all(cell is None or str(cell).strip() == "" for cell in raw):
            continue
        qid = str(raw[mapping["question_id"]]).strip()
        if qid not in EXPECTED_IDS:
            raise AnalysisError(f"Unexpected question id in human review: {qid!r}")
        if qid in parsed:
            raise AnalysisError(f"Duplicate human review row: {qid}")
        parsed[qid] = {
            name: parse_score(raw[mapping[name]]) for name in HUMAN_DIMENSIONS
        }
    validate_id_set(parsed.keys(), "human review")
    meta = {"sheet": sheet_name, "headers": headers, "column_mapping": mapping}
    return parsed, meta


def require_validated_archive_scores(human: dict[str, dict[str, Any]]) -> None:
    """Reject a review sheet that does not contain the five confirmed corrections."""
    mismatches: list[str] = []
    for question_id, expected in CONFIRMED_ARCHIVE_SCORES.items():
        actual = human[question_id]["human_archive_correctness"]
        if actual != expected:
            mismatches.append(f"{question_id}={actual!r}, expected {expected}")
    if mismatches:
        raise AnalysisError(
            "Human review is not the validated workbook. "
            "human_archive_correctness mismatches: " + "; ".join(mismatches)
        )


def validate_id_set(ids: Iterable[str], label: str) -> None:
    found = list(ids)
    if found != EXPECTED_IDS and sorted(found) != EXPECTED_IDS:
        missing = [qid for qid in EXPECTED_IDS if qid not in found]
        extra = [qid for qid in found if qid not in EXPECTED_IDS]
        dupes = [qid for qid in EXPECTED_IDS if found.count(qid) > 1]
        raise AnalysisError(
            f"{label} ids mismatch missing={missing} extra={extra} duplicates={dupes}"
        )
    if len(found) != len(EXPECTED_IDS):
        raise AnalysisError(f"{label} does not contain exactly 28 rows")


def index_phase1(rows: list[dict[str, Any]]) -> dict[str, dict[str, dict[str, Any]]]:
    grouped: dict[str, dict[str, dict[str, Any]]] = {cfg: {} for cfg in CONFIGS}
    for row in rows:
        cfg = str(row["config_id"])
        qid = str(row["question_id"])
        if cfg not in grouped:
            raise AnalysisError(f"Unexpected config_id: {cfg}")
        if qid in grouped[cfg]:
            raise AnalysisError(f"Duplicate Phase 1 row {cfg} {qid}")
        grouped[cfg][qid] = row
    for cfg, by_id in grouped.items():
        validate_id_set(by_id.keys(), f"phase1 {cfg}")
    return grouped


def verify_stored_ir_metrics(grouped: dict[str, dict[str, dict[str, Any]]]) -> dict[str, Any]:
    """Confirm stored IR metrics match score_ranking, including P@5 denominator k=5."""
    mismatches: list[str] = []
    short_rankings = 0
    for cfg, by_id in grouped.items():
        for qid, row in by_id.items():
            scored = score_ranking(
                row["ranked_pfc_ids"],
                row["primary_relevant"],
                unsupported=bool(row["unsupported"]),
            )
            for metric in IR_METRICS:
                stored = row[metric]
                recomputed = scored[metric]
                if stored is None and recomputed is None:
                    continue
                if stored is None or recomputed is None or not math.isclose(
                    float(stored), float(recomputed), rel_tol=0, abs_tol=1e-12
                ):
                    mismatches.append(f"{cfg}:{qid}:{metric}")
            if not row["unsupported"] and len(row["ranked_pfc_ids"]) < 5:
                short_rankings += 1
                relevant = {str(item) for item in row["primary_relevant"]}
                hits = sum(1 for doc in row["ranked_pfc_ids"][:5] if doc in relevant)
                expected = hits / 5.0
                if not math.isclose(float(row["precision_at_5"]), expected, abs_tol=1e-12):
                    mismatches.append(f"{cfg}:{qid}:precision_denominator")
    return {
        "mismatches": mismatches,
        "supported_rankings_shorter_than_5": short_rankings,
        "precision_at_5_rule": (
            "precision_at_k divides the number of relevant documents in the top min(k, n) "
            "by k (k=5), not by the number of documents returned. An empty ranking is 0. "
            "This is the rule already stored in retrieval.jsonl."
        ),
    }


def _supported(by_id: dict[str, dict[str, Any]]) -> list[dict[str, Any]]:
    return [by_id[qid] for qid in EXPECTED_IDS if not by_id[qid]["unsupported"]]


def phase1_config_rows(
    grouped: dict[str, dict[str, dict[str, Any]]],
) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for cfg in CONFIGS:
        supported = _supported(grouped[cfg])
        unsupported_n = sum(1 for row in grouped[cfg].values() if row["unsupported"])
        for metric in IR_METRICS:
            values = [row[metric] for row in supported]
            stats = describe(values)
            ci = bootstrap_mean_ci(values, seed_key=f"phase1:{cfg}:{metric}")
            rows.append(
                {
                    "config_id": cfg,
                    "subset": "supported",
                    "category": "",
                    "metric": metric,
                    "supported_n": len(supported),
                    "unsupported_n": unsupported_n,
                    **stats,
                    "ci95_low": ci["ci95_low"],
                    "ci95_high": ci["ci95_high"],
                    "p90": None,
                    "p95": None,
                    "range": None,
                }
            )
        latency = latency_describe(row["retrieval_duration_ms"] for row in supported)
        ci = bootstrap_mean_ci(
            [row["retrieval_duration_ms"] for row in supported],
            seed_key=f"phase1:{cfg}:retrieval_duration_ms",
        )
        rows.append(
            {
                "config_id": cfg,
                "subset": "supported",
                "category": "",
                "metric": "retrieval_duration_ms",
                "supported_n": len(supported),
                "unsupported_n": unsupported_n,
                "n": latency["n"],
                "mean": latency["mean"],
                "std": latency["std"],
                "median": latency["median"],
                "min": latency["min"],
                "max": latency["max"],
                "ci95_low": ci["ci95_low"],
                "ci95_high": ci["ci95_high"],
                "p90": latency["p90"],
                "p95": latency["p95"],
                "range": latency["range"],
            }
        )
        categories = sorted({row["category"] for row in supported})
        for category in categories:
            members = [row for row in supported if row["category"] == category]
            for metric in IR_METRICS:
                stats = describe(row[metric] for row in members)
                rows.append(
                    {
                        "config_id": cfg,
                        "subset": "supported_category",
                        "category": category,
                        "metric": metric,
                        "supported_n": len(supported),
                        "unsupported_n": unsupported_n,
                        **stats,
                        "ci95_low": None,
                        "ci95_high": None,
                        "p90": None,
                        "p95": None,
                        "range": None,
                    }
                )
    return rows


def _aligned(grouped, left: str, right: str, field: str) -> tuple[list[str], np.ndarray]:
    ids = [qid for qid in EXPECTED_IDS if not grouped[left][qid]["unsupported"]]
    diffs = []
    for qid in ids:
        a = grouped[left][qid][field]
        b = grouped[right][qid][field]
        if a is None or b is None:
            raise AnalysisError(f"Missing paired value {field} {left}/{right} {qid}")
        diffs.append(float(a) - float(b))
    return ids, np.asarray(diffs, dtype=float)


def phase1_paired(
    grouped: dict[str, dict[str, dict[str, Any]]],
) -> tuple[list[dict[str, Any]], list[dict[str, Any]], dict[str, Any]]:
    comparison_rows: list[dict[str, Any]] = []
    query_rows: list[dict[str, Any]] = []
    tests: list[tuple[str, float]] = []
    for baseline, alternative in COMPARISONS:
        # Difference is alternative minus baseline (second minus first in the requested pair).
        ids, _ = _aligned(grouped, alternative, baseline, "mrr")
        for metric in PAIRED_METRICS:
            _, diffs = _aligned(grouped, alternative, baseline, metric)
            stats = describe(diffs.tolist())
            ci = bootstrap_mean_ci(
                diffs.tolist(), seed_key=f"paired:{alternative}-{baseline}:{metric}"
            )
            improved = [qid for qid, diff in zip(ids, diffs, strict=True) if diff > 0]
            worsened = [qid for qid, diff in zip(ids, diffs, strict=True) if diff < 0]
            tied = [qid for qid, diff in zip(ids, diffs, strict=True) if diff == 0]
            test = wilcoxon_paired(diffs.tolist())
            key = f"{alternative}_minus_{baseline}:{metric}"
            if test["pvalue"] is not None:
                tests.append((key, test["pvalue"]))
            comparison_rows.append(
                {
                    "comparison": f"{alternative}_minus_{baseline}",
                    "baseline": baseline,
                    "alternative": alternative,
                    "metric": metric,
                    "difference": f"{alternative}_minus_{baseline}",
                    "n": stats["n"],
                    "mean_difference": stats["mean"],
                    "std_difference": stats["std"],
                    "median_difference": stats["median"],
                    "min_difference": stats["min"],
                    "max_difference": stats["max"],
                    "ci95_low": ci["ci95_low"],
                    "ci95_high": ci["ci95_high"],
                    "n_improved": len(improved),
                    "n_worsened": len(worsened),
                    "n_tied": len(tied),
                    "improved_ids": " ".join(improved),
                    "worsened_ids": " ".join(worsened),
                    "wilcoxon_applicable": test["applicable"],
                    "wilcoxon_statistic": test["statistic"],
                    "pvalue_unadjusted": test["pvalue"],
                    "n_ties_excluded": test["n_ties"],
                    "n_nonzero": test["n_nonzero"],
                    "wilcoxon_reason": test["reason"] or "",
                    "pvalue_holm": None,
                }
            )
        _, latency = _aligned(grouped, alternative, baseline, "retrieval_duration_ms")
        lat_stats = describe(latency.tolist())
        lat_ci = bootstrap_mean_ci(
            latency.tolist(),
            seed_key=f"paired:{alternative}-{baseline}:retrieval_duration_ms",
        )
        comparison_rows.append(
            {
                "comparison": f"{alternative}_minus_{baseline}",
                "baseline": baseline,
                "alternative": alternative,
                "metric": "retrieval_duration_ms",
                "difference": f"{alternative}_minus_{baseline}",
                "n": lat_stats["n"],
                "mean_difference": lat_stats["mean"],
                "std_difference": lat_stats["std"],
                "median_difference": lat_stats["median"],
                "min_difference": lat_stats["min"],
                "max_difference": lat_stats["max"],
                "ci95_low": lat_ci["ci95_low"],
                "ci95_high": lat_ci["ci95_high"],
                "n_improved": "",
                "n_worsened": "",
                "n_tied": int(np.sum(latency == 0)),
                "improved_ids": "",
                "worsened_ids": "",
                "wilcoxon_applicable": False,
                "wilcoxon_statistic": None,
                "pvalue_unadjusted": None,
                "n_ties_excluded": None,
                "n_nonzero": None,
                "wilcoxon_reason": "latency_reported_descriptively_not_in_holm_family",
                "pvalue_holm": None,
            }
        )
        for qid in ids:
            query_rows.append(
                {
                    "comparison": f"{alternative}_minus_{baseline}",
                    "question_id": qid,
                    "delta_mrr": grouped[alternative][qid]["mrr"] - grouped[baseline][qid]["mrr"],
                    "delta_recall_at_5": grouped[alternative][qid]["recall_at_5"]
                    - grouped[baseline][qid]["recall_at_5"],
                    "delta_recall_at_10": grouped[alternative][qid]["recall_at_10"]
                    - grouped[baseline][qid]["recall_at_10"],
                    "delta_precision_at_5": grouped[alternative][qid]["precision_at_5"]
                    - grouped[baseline][qid]["precision_at_5"],
                    "delta_retrieval_duration_ms": grouped[alternative][qid][
                        "retrieval_duration_ms"
                    ]
                    - grouped[baseline][qid]["retrieval_duration_ms"],
                }
            )
    adjusted = holm_adjust(tests)
    for row in comparison_rows:
        key = f"{row['comparison']}:{row['metric']}"
        if key in adjusted:
            row["pvalue_holm"] = adjusted[key]
    identity = _config_identity(grouped, "A3", "A4")
    mrr_focus = _mrr_delta_focus(grouped, "A5", "A3", ("Q009", "Q011"))
    return comparison_rows, query_rows, {"a3_vs_a4": identity, "a5_vs_a3_mrr": mrr_focus, "holm_family_size": len(tests)}


def _config_identity(grouped, left: str, right: str) -> dict[str, Any]:
    metric_identical = True
    ranking_identical = True
    differing = []
    for qid in EXPECTED_IDS:
        a = grouped[left][qid]
        b = grouped[right][qid]
        if a["ranked_pfc_ids"] != b["ranked_pfc_ids"]:
            ranking_identical = False
            differing.append(qid)
        for metric in (*IR_METRICS,):
            if a[metric] != b[metric]:
                metric_identical = False
                if qid not in differing:
                    differing.append(qid)
    return {
        "metric_identical_all_28": metric_identical,
        "ranking_identical_all_28": ranking_identical,
        "differing_question_ids": differing,
    }


def _mrr_delta_focus(grouped, left: str, right: str, focus: Sequence[str]) -> dict[str, Any]:
    ids, diffs = _aligned(grouped, left, right, "mrr")
    changed = {
        qid: float(diff)
        for qid, diff in zip(ids, diffs, strict=True)
        if diff != 0
    }
    mean_delta = float(diffs.mean()) if diffs.size else None
    focus_sum = sum(changed.get(qid, 0.0) for qid in focus)
    other = {qid: value for qid, value in changed.items() if qid not in focus}
    return {
        "mean_mrr_difference": mean_delta,
        "changed_question_ids": changed,
        "focus_ids": list(focus),
        "sum_of_focus_differences": focus_sum,
        "other_changed": other,
        "focus_accounts_for_entire_sum": not other and set(changed) <= set(focus),
    }


def write_csv(path: Path, rows: list[dict[str, Any]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    if not rows:
        path.write_text("", encoding="utf-8")
        return
    fieldnames: list[str] = []
    for row in rows:
        for key in row:
            if key not in fieldnames:
                fieldnames.append(key)
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames, extrasaction="ignore")
        writer.writeheader()
        for row in rows:
            writer.writerow({key: _csv_value(row.get(key)) for key in fieldnames})


def _csv_value(value: Any) -> Any:
    if value is None:
        return ""
    if isinstance(value, float):
        return format(value, ".12g")
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False)
    return value
